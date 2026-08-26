from pathlib import Path
import json
import openpyxl
import sys

sys.stdout.reconfigure(encoding="utf-8", errors="replace")


ROOT = Path(__file__).resolve().parents[2]
WORKBOOKS = [
    ROOT / "public" / "GEAPP Reporting.xlsx",
    ROOT / "public" / "M 300 Initial Setup Questions.xlsx",
    ROOT.parent / "m300-demo-backend" / "EMRC - M300 Data Framework v0-5 2026-06-4.xlsx",
]


def cell_value(cell):
    value = cell.value
    if value is None:
        return None
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return str(value)


for path in WORKBOOKS:
    workbook = openpyxl.load_workbook(path, data_only=False, read_only=False)
    print(f"\nWORKBOOK: {path}")
    print("SHEETS:", json.dumps(workbook.sheetnames, ensure_ascii=False))
    for sheet in workbook.worksheets:
        print(f"\nSHEET: {sheet.title} rows={sheet.max_row} cols={sheet.max_column}")
        emitted = 0
        for row in sheet.iter_rows():
            values = [cell_value(cell) for cell in row]
            if not any(value not in (None, "") for value in values):
                continue
            while values and values[-1] in (None, ""):
                values.pop()
            print(f"R{row[0].row}: {json.dumps(values, ensure_ascii=False)}")
            emitted += 1
            if emitted >= 45:
                remaining = sheet.max_row - row[0].row
                if remaining > 0:
                    print(f"... {remaining} later worksheet rows not shown")
                break


pdf_path = ROOT / "public" / "Snapshot of CDMU Convening Slides.pptx.pdf"
print(f"\nPDF: {pdf_path}")
try:
    from pypdf import PdfReader
    reader = PdfReader(pdf_path)
    print(f"PAGES: {len(reader.pages)}")
    for index, page in enumerate(reader.pages, 1):
        text = " ".join((page.extract_text() or "").split())
        print(f"P{index}: {text[:700]}")
except Exception as error:
    print(f"PDF_EXTRACTION_ERROR: {error}")
