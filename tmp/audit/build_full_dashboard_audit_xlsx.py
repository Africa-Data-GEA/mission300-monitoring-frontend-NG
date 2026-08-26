from pathlib import Path
from datetime import date

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo


ROOT = Path(__file__).resolve().parents[2]
OUTPUT_DIR = ROOT / "outputs" / "01a032df-c6f2-7622-ab51-3bd4d3351c18"
OUTPUT_PATH = OUTPUT_DIR / "M300_Full_Dashboard_Backend_Logic_Audit.xlsx"

NAVY = "17365D"
BLUE = "1F4E78"
TEAL = "0F6B78"
LIGHT_BLUE = "D9EAF7"
PALE_BLUE = "EAF3F8"
LIGHT_GREY = "F2F4F7"
MID_GREY = "D0D5DD"
DARK = "1D2939"
WHITE = "FFFFFF"
RED = "B42318"
RED_FILL = "FEE4E2"
ORANGE = "B54708"
ORANGE_FILL = "FEF0C7"
YELLOW = "8A6100"
YELLOW_FILL = "FFF4CC"
GREEN = "027A48"
GREEN_FILL = "D1FADF"
PURPLE = "6941C6"
PURPLE_FILL = "F4EBFF"

THIN = Side(style="thin", color=MID_GREY)
TABLE_BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def title_band(ws, title, subtitle, last_col):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    cell = ws.cell(1, 1, title)
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.font = Font(name="Aptos Display", size=18, bold=True, color=WHITE)
    cell.alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 34
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)
    cell = ws.cell(2, 1, subtitle)
    cell.fill = PatternFill("solid", fgColor=PALE_BLUE)
    cell.font = Font(name="Aptos", size=10, italic=True, color=DARK)
    cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[2].height = 28


def style_header(row):
    for cell in row:
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.font = Font(name="Aptos", size=10, bold=True, color=WHITE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = TABLE_BORDER


def style_body(ws, start_row, end_row, start_col, end_col, heights=55):
    for row in ws.iter_rows(min_row=start_row, max_row=end_row, min_col=start_col, max_col=end_col):
        for cell in row:
            cell.font = Font(name="Aptos", size=9, color=DARK)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = TABLE_BORDER
        ws.row_dimensions[row[0].row].height = heights


def add_table(ws, name, ref):
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
        showRowStripes=True, showColumnStripes=False
    )
    ws.add_table(table)


def add_status_validation(ws, cell_range):
    validation = DataValidation(
        type="list",
        formula1='"Backend Response Required,Open,In Progress,Resolved,Accepted Risk"',
        allow_blank=True,
    )
    validation.error = "Select a status from the dropdown."
    validation.errorTitle = "Invalid status"
    ws.add_data_validation(validation)
    validation.add(cell_range)


def add_priority_validation(ws, cell_range):
    validation = DataValidation(type="list", formula1='"P0,P1,P2,P3"', allow_blank=True)
    ws.add_data_validation(validation)
    validation.add(cell_range)


def set_print(ws, landscape=True, repeat_rows=None):
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = 0.25
    ws.page_margins.right = 0.25
    ws.page_margins.top = 0.4
    ws.page_margins.bottom = 0.4
    if repeat_rows:
        ws.print_title_rows = repeat_rows


modules = [
    ["F-001", "Authentication", "High", "Partial", "Secure account lifecycle, real reset/invite completion, role and scope context, revocation and security audit.", "Database login and lockout exist, but forgot-password is a stub, invitations send no email, access scopes are unused, and permission context is incomplete.", "Invited users cannot complete onboarding; access decisions cannot be made reliably.", "Implement reset/invite lifecycle, token revocation/logout, complete role/scope payload and security-event audit.", "P1"],
    ["F-002", "Authorization & Access Scope", "Critical", "Critical failure", "Server-enforced least privilege for View, Submit, Edit, Validate, Approve, Export and Administer, including institution/dataset/geography/programme scope.", "Most endpoints require only an authenticated token. Role checks exist on only a few writes; frontend navigation hiding is acting as the main restriction.", "A provider can access Administration and unrelated institutional records. This is a confidentiality and integrity risk.", "Create and test an endpoint-by-endpoint permission and record-scope policy; deny out-of-scope access server-side.", "P0"],
    ["F-003", "Executive Overview", "Critical", "Placeholder", "Calculated snapshot from validated KPI, target and implementation records; every filter applied with traceable provenance and current freshness.", "Headline, access, performance and project values are fixed controller arrays. Period/technology filters are ignored and other filters affect only parts of the payload.", "The page appears live while showing non-calculated, stale or unfiltered values.", "Rebuild the response from accepted observations and implementation records; add filter correctness and empty-slice tests.", "P1"],
    ["F-004", "Pillar Dashboards", "Critical", "Mostly placeholder", "Each pillar should filter real KPI, programme/project and bottleneck records and provide useful record-level drill-downs.", "Only parts of the KPI list are database-backed. Headline/delivery values are representative placeholders and bottlenecks are empty because joins are missing.", "Different pillars can display repeated generic information with no operational drill-down.", "Build pillar queries over the shared KPI/implementation/bottleneck data model and implement record relationships.", "P1"],
    ["F-005", "KPI Explorer", "High", "Partial catalogue", "Governed KPI master with formula, mappings, targets, source, readiness, history, versioning, lineage and add/edit/deactivate workflow.", "Some metadata and observation tables are read, but PHP has no general calculation/publication engine, incomplete metadata dimensions and no add/deactivate workflow.", "KPI values cannot be proven to originate from approved submissions, and catalogue governance is incomplete.", "Complete KPI metadata/version model and link published observations transactionally to accepted submissions or approved public sources.", "P1"],
    ["F-006", "Programmes & Projects", "High", "Incomplete/incorrect", "Full programme, project, action and milestone lifecycle with dates, owners, evidence, history, filters and editing.", "Reads exist, but create/update/delete routes used by the UI are missing. Institution filtering uses the wrong field. On-time logic does not compare actual and due dates; pending/missed are mislabeled.", "Users see incorrect delivery performance and cannot complete the implementation workflow.", "Implement authorized CRUD, actions/evidence/history, correct relationship filters, and approved status/date formulas.", "P2"],
    ["F-007", "Bottlenecks & Exceptions", "High", "Incomplete workflow", "Linked register with severity, lifecycle, accountable owner, age, escalation, decisions, resolution, evidence, history and audit.", "Create/list/by-project are partial; frontend and backend creation fields differ. Update, escalation, decision, resolution and evidence workflows are absent.", "Constraints cannot be managed through closure or reliably tied to delivery records.", "Align the request contract and implement full lifecycle transitions, relationships, authorization, evidence and audit.", "P2"],
    ["F-008", "Data Submission & Validation", "Critical", "Simulation", "Versioned templates; preview; parse; field/business validation; calculation; status history; scoped review; publication to KPI history; compliance and audit.", "Obligations/fields are fixed; files are not parsed; validation and KPI values are hardcoded; temporary storage is used; decisions do not publish to PHP KPI tables.", "Uploaded data does not actually drive the dashboard even when the endpoint succeeds.", "Implement the demonstrated parse → validate → calculate → decide → publish → audit transaction using the PHP database schema.", "P1"],
    ["F-009", "State & DisCo Performance", "Critical", "Mock data", "Validated metering, ATC&C, remittance, tariff, supply and delivery calculations filtered consistently by period, DisCo, state and source.", "Fixed MockDb arrays cover only a small sample and are disconnected from submissions and observations.", "Utility and state comparisons are not evidence-based and may fabricate coverage.", "Aggregate only accepted source records using approved weighted formulas; never invent unavailable state/LGA values.", "P1"],
    ["F-010", "Reports & Exports", "High", "Mock generator", "Custom and standard reports generated from approved live records, with saved configurations, CPR mapping, real files, versions and audit.", "Catalogue/preview rows are fixed; saved reports use temporary storage; live modules are not queried; file references are not genuine and methodology route is missing.", "Exports cannot be reconciled to the dashboard or used safely for reporting obligations.", "Generate reports from the same approved observations/register records; persist runs/files/configurations and provide complete provenance.", "P2"],
    ["F-011", "Stakeholders & Institutions", "Medium", "Partial directory", "Institution registry with custodians, focal roles, datasets, programmes/projects, geographies and responsibility assignments.", "Some registry and custodian data is database-backed, but the full responsibility/scope model is absent and reads are broadly accessible.", "Accountability and assignment workflows cannot be enforced or reported completely.", "Complete institution-to-dataset/programme/project/geography mappings and apply them to authorization and queues.", "P2"],
    ["F-012", "User Administration", "Critical", "Incomplete/overexposed", "Complete lifecycle, role/permission matrix, multiple access scopes, edit/suspend/reactivate, MFA policy, invitations and audit.", "User listing/invite creation exist, but invitation delivery/reset completion and lifecycle operations are missing. Administration reads are not admin-only.", "User onboarding is broken and sensitive administrative data is exposed.", "Restrict all administration operations, complete onboarding and account lifecycle, manage scopes and record every change.", "P0"],
    ["F-013", "Public Gateway", "High", "Integration failure", "Public, safe, approved indicators with provenance through an unauthenticated public API.", "PHP provides a fixed public endpoint, but the frontend public pages call protected executive/state endpoints.", "Public pages may fail without login and do not have a controlled publication boundary.", "Make the frontend consume the public endpoint and return only explicitly publishable approved observations.", "P1"],
    ["F-014", "Branding", "Medium", "Partial", "Persisted country configuration with admin-only writes and a functional reset.", "Writes are admin-gated but use file/mock persistence; the frontend reset route does not exist in PHP.", "Configuration can drift and reset fails in live backend mode.", "Use governed persistence and implement the contracted reset endpoint with audit.", "P2"],
    ["F-015", "Audit & Activity", "Critical", "Missing", "Immutable cross-module events for login, upload, edit, validation, approval, export and administration.", "An audit table exists, but controllers do not write platform events and no activity API is available.", "Changes and decisions cannot be reconstructed or defended.", "Create a shared audit service, transactional event writes, retention policy and scoped audit query API.", "P0"],
]

missing_endpoints = [
    ["E-001", "GET", "/api/data-submissions/datasets", "Data Submission", "Dataset catalogue behind obligations, templates, fields, ownership and KPI mappings.", "Frontend request fails in backend mode.", "P0"],
    ["E-002", "GET", "/api/reports/methodology", "Reports", "Definitions, formulae, sources and limitations used in generated reports.", "Reports methodology section cannot load.", "P1"],
    ["E-003", "POST", "/api/branding/reset", "Branding", "Restore approved default branding configuration.", "Reset action fails in backend mode.", "P2"],
    ["E-004", "POST", "/api/programs", "Implementation Register", "Create an authorized, validated programme with history and audit.", "Add Programme page cannot complete.", "P1"],
    ["E-005", "POST", "/api/programs/projects", "Implementation Register", "Create a project linked to programme, institution, pillar and location.", "Add Project page cannot complete.", "P1"],
    ["E-006", "PATCH", "/api/programs/projects/:projectId", "Implementation Register", "Update project fields with version/history and authorization.", "Edit Project page cannot save.", "P1"],
    ["E-007", "DELETE", "/api/programs/projects/:projectId", "Implementation Register", "Authorized archive/delete behavior with dependency checks and audit.", "Delete action cannot complete.", "P2"],
]

actions = [
    ["A-001", "P0", "Inventory every fixed, MockDb, seed-only and temporary-store response; label or disable it until approved lineage exists.", "All modules", "No mock response is presented as confirmed/live; inventory is signed off.", "Backend Lead"],
    ["A-002", "P0", "Enforce permission and record scope on every endpoint.", "Security", "Automated role/scope matrix tests reject every unauthorized case.", "Security Lead"],
    ["A-003", "P0", "Remove provider access to Administration and unrelated institutional obligations.", "Security", "Provider receives 403/filtered results in live verification.", "Security Lead"],
    ["A-004", "P0", "Agree one API contract, ID strategy and controlled vocabulary.", "Architecture", "Frontend schema contract tests pass for all routes.", "API Lead"],
    ["A-005", "P0", "Confirm the authoritative Nigeria backend, database and environments.", "Governance", "Written environment/data-source statement is provided.", "Product Owner"],
    ["A-006", "P1", "Complete pillar → programme → project → action/milestone → evidence/bottleneck relationships.", "Data model", "Every drill-down resolves through stable database relationships.", "Data Lead"],
    ["A-007", "P1", "Complete institution, user, role and multi-scope assignments.", "Access control", "Users can hold multiple effective scoped assignments.", "Security Lead"],
    ["A-008", "P1", "Complete dataset, field, controlled reference, template-version, formula and KPI mappings.", "Data framework", "Each KPI has approved, versioned input-field dependencies.", "Data Lead"],
    ["A-009", "P1", "Implement database-backed parse → validate → calculate → decide → publish workflow.", "Submissions", "End-to-end tests prove approved uploads update KPI history and obligations.", "Backend Lead"],
    ["A-010", "P1", "Implement immutable observation lineage and cross-module audit events.", "Audit", "Every displayed value/action links to source and actor/time.", "Backend Lead"],
    ["A-011", "P1", "Rebuild Overview, Pillars and State/DisCo from accepted observations and implementation records.", "Dashboards", "No controller-level metric arrays; filter/empty-slice tests pass.", "API Lead"],
    ["A-012", "P2", "Implement programme/project/action CRUD with evidence and history.", "Implementation Register", "All frontend operations work with authorization and audit.", "Backend Lead"],
    ["A-013", "P2", "Implement bottleneck escalation, decision, resolution and follow-up.", "Bottlenecks", "Complete governed lifecycle is demonstrated end-to-end.", "Backend Lead"],
    ["A-014", "P2", "Complete invitation/reset, suspension/reactivation and access-scope administration.", "Administration", "Invited user can activate; admin can manage lifecycle securely.", "Security Lead"],
    ["A-015", "P2", "Generate reports from live approved records and support saved custom configurations.", "Reports", "Report values reconcile exactly to dashboard records.", "Reporting Lead"],
    ["A-016", "P2", "Connect public pages exclusively to the safe public API.", "Public Gateway", "Public views load without authentication and expose only publishable data.", "API Lead"],
    ["A-017", "P2", "Implement KPI add, edit, deactivate and version workflow.", "KPI Explorer", "Governed KPI lifecycle is available and audited.", "Data Lead"],
    ["A-018", "P3", "Add unit tests for each approved formula and boundary condition.", "Quality", "Formula test pack passes and reconciles to approved examples.", "QA Lead"],
    ["A-019", "P3", "Add authorization and scope integration tests.", "Quality", "All allow/deny combinations in the approved matrix pass.", "QA Lead"],
    ["A-020", "P3", "Add response contract tests generated from frontend schemas.", "Quality", "No runtime schema mismatch across all endpoints.", "QA Lead"],
    ["A-021", "P3", "Add submission-to-dashboard/report end-to-end tests.", "Quality", "One accepted submission updates every dependent view consistently.", "QA Lead"],
    ["A-022", "P3", "Add filter correctness and non-matching empty-slice tests.", "Quality", "Every supplied filter changes the server data slice correctly.", "QA Lead"],
    ["A-023", "P3", "Reconcile dashboard values against approved workbooks/publications.", "Data quality", "Signed reconciliation report has no unexplained variance.", "Data Lead"],
    ["A-024", "P3", "Complete UAT with CDMU, SEforALL and designated custodians.", "Acceptance", "Formal UAT sign-off and defect closure are recorded.", "Product Owner"],
]

questions = [
    [1, "Is https://nigeria.geappindia.org the authoritative Nigeria production API and database, or is there another environment containing the real data?", "Environment & data ownership"],
    [2, "Which modules are declared production-ready, and which are acknowledged mock or representative implementations?", "Implementation status"],
    [3, "For every displayed KPI, which database observation, accepted submission, source reference and formula produced the value?", "KPI lineage"],
    [4, "Where is the approved formula-to-dataset-field mapping requested in the initial setup workbook?", "Calculation governance"],
    [5, "How are validation decisions published transactionally to KPI observations and obligations?", "Submission workflow"],
    [6, "How are programme, project, action, milestone and bottleneck relationships represented and maintained?", "Implementation data model"],
    [7, "What server-side permission and record-scope rule protects each endpoint?", "Authorization"],
    [8, "Why can a provider account access Administration data and unrelated obligations?", "Security finding"],
    [9, "Which status vocabularies and transition diagrams have been approved?", "Workflow governance"],
    [10, "Which audit events are written for login, upload, edit, validation, approval, export and administration?", "Audit"],
    [11, "How are reports generated from live approved records, and where are generated files and version histories stored?", "Reporting"],
    [12, "How will the public gateway receive approved public data without authentication?", "Public publication"],
    [13, "What is the migration plan from fixed arrays, MockDb and temporary-store responses to database-backed logic?", "Remediation"],
    [14, "Which acceptance tests demonstrate calculation accuracy, filter behavior, authorization and end-to-end publication?", "Quality assurance"],
]

criteria = [
    ["G-001", "Authoritative source", "Displayed records originate from an identified and approved source.", "Provide source-to-record trace for sampled values.", "Not Tested"],
    ["G-002", "Formula accuracy", "Formulae and field mappings are approved, versioned and tested.", "Reconcile test cases to approved workbook examples.", "Not Tested"],
    ["G-003", "Metric metadata", "Every value has period, unit, target basis, validation state and provenance.", "Contract/data-quality scan has zero omissions.", "Not Tested"],
    ["G-004", "Filter behavior", "Every filter changes the server-side slice and no-match queries return empty sections.", "Automated positive and empty-slice tests pass.", "Not Tested"],
    ["G-005", "Workflow completion", "Authorized users can complete the intended workflow from start to finish.", "End-to-end workflow tests pass.", "Not Tested"],
    ["G-006", "Access control", "Unauthorized and out-of-scope users are rejected by the API.", "Role/scope matrix tests pass with 403 or scoped results.", "Failed"],
    ["G-007", "State consistency", "Every transition persists and appears consistently across modules.", "Cross-page reconciliation after each transition passes.", "Not Tested"],
    ["G-008", "Auditability", "Important actions create immutable, queryable audit events.", "Sample login/upload/edit/decision/export/admin events exist.", "Failed"],
    ["G-009", "Report reconciliation", "Reports use and reconcile to the same approved records shown in the dashboard.", "Sample report matches source views and versions exactly.", "Failed"],
    ["G-010", "No mock dependency", "Production paths operate without frontend mocks, MockDb or fixed controller responses.", "Backend-mode end-to-end test passes with mock systems disabled.", "Failed"],
]

sources = [
    ["S-001", "Approved reporting workbook", "public/GEAPP Reporting.xlsx", "Action/target statuses, comments, update dates, source rules and Nigeria reference values."],
    ["S-002", "Backend setup questions", "public/M 300 Initial Setup Questions.xlsx", "Hierarchy, roles, feature permissions, dataset mappings, formula requirements and activity APIs."],
    ["S-003", "CDMU framework slides", "public/Snapshot of CDMU Convening Slides.pptx.pdf", "Monitoring, reporting and learning; coordination, bottlenecks, evidence and delivery outcomes."],
    ["S-004", "Inception report", "docs/EMRC - Inception Report for the Design and Implementation of the M300 Monitoring Dashboard for Nigeria v0-1 2026-02-16.docx", "Monitoring architecture, governance, validation, reporting, traceability and access requirements."],
    ["S-005", "KPI matrix", "docs/EMRC - M300 KPI Matrix v2-0 2026-06-29.xlsx", "Definitions, sources, baselines, targets, validation/confidence states and known gaps."],
    ["S-006", "Platform feedback", "docs/M300 Monitoring Platform Feedback.pptx", "Dynamic KPI data, filters/sorts, drill-down tables, linked records, preview, custom reports and KPI lifecycle."],
    ["S-007", "Page specifications", "docs/source/EMRC - M300 Dashboard Page Specifications v0-3 2026-07-17.xlsx", "Expected elements, formulas, targets, filters, methodologies and sources for all modules."],
    ["S-008", "Wireframe", "docs/source/EMRC - M300 Dashboard Configuration Wireframe v0-1 2026-06-10.pdf", "Configuration and dashboard behavior reference."],
    ["S-009", "Frontend API contract", "docs/API.md", "Runtime response schemas, routes, filter behavior and workflow contracts."],
    ["S-010", "Module specifications", "docs/specs/", "Detailed expected behavior for every dashboard module."],
    ["S-011", "Frontend source", "src/", "Actual routes, API calls, schemas, pages and backend-mode behavior."],
    ["S-012", "PHP route table", "mission300-monitoring-backend-NG/api/routes.php", "Implemented deployed-backend route inventory."],
    ["S-013", "PHP controllers", "mission300-monitoring-backend-NG/api/src/App/Controllers/", "Actual calculations, filters, mock arrays, persistence and authorization calls."],
    ["S-014", "PHP database schema", "mission300-monitoring-backend-NG/api/nigeria_m300_mysql.sql", "Defined entities compared with controller usage."],
    ["S-015", "Demo backend contract", "../m300-demo-backend/docs/BACKEND-CONTRACT.md", "Reference parse/validate/calculate/review/publish/audit design and stated scope limitations."],
    ["S-016", "Demo backend tests/source", "../m300-demo-backend/src and test/", "Two-dataset calculation pipeline and end-to-end verification."],
]


wb = Workbook()
summary = wb.active
summary.title = "Executive Summary"

findings = wb.create_sheet("Module Findings")
endpoints = wb.create_sheet("Missing Endpoints")
priority = wb.create_sheet("Priority Actions")
questions_ws = wb.create_sheet("Backend Questions")
criteria_ws = wb.create_sheet("Go-Live Criteria")
sources_ws = wb.create_sheet("Evidence Sources")

# Executive Summary
title_band(summary, "M300 Full Dashboard Backend Logic Audit", "Prepared for backend-team review | Scope: entire platform | No backend code or data changed", 10)
summary.merge_cells("A4:F4")
summary["A4"] = "Overall conclusion"
summary["A4"].fill = PatternFill("solid", fgColor=TEAL)
summary["A4"].font = Font(name="Aptos", size=12, bold=True, color=WHITE)
summary.merge_cells("A5:F8")
summary["A5"] = (
    "The deployed PHP backend is not production-ready for the dashboard represented by the frontend. "
    "It combines database reads with fixed arrays, MockDb data and temporary persistence. Successful HTTP responses "
    "do not prove that values are calculated, validated, scoped, auditable or connected across workflows."
)
summary["A5"].alignment = Alignment(vertical="top", wrap_text=True)
summary["A5"].font = Font(name="Aptos", size=11, color=DARK)
summary["A5"].fill = PatternFill("solid", fgColor=PALE_BLUE)
summary["A5"].border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

summary["A10"] = "Finding summary"
summary["A10"].font = Font(name="Aptos", size=12, bold=True, color=NAVY)
summary["A11"] = "Severity"
summary["B11"] = "Count"
style_header(summary[11][:2])
for idx, severity in enumerate(["Critical", "High", "Medium"], start=12):
    summary.cell(idx, 1, severity)
    summary.cell(idx, 2, f'=COUNTIF(\'Module Findings\'!$C$5:$C$19,A{idx})')
    summary.cell(idx, 1).border = TABLE_BORDER
    summary.cell(idx, 2).border = TABLE_BORDER
    summary.cell(idx, 2).alignment = Alignment(horizontal="center")

summary["D10"] = "Audit facts"
summary["D10"].font = Font(name="Aptos", size=12, bold=True, color=NAVY)
facts = [
    ["Modules assessed", "=COUNTA('Module Findings'!$B$5:$B$19)"],
    ["Missing frontend routes", "=COUNTA('Missing Endpoints'!$C$5:$C$11)"],
    ["Backend actions", "=COUNTA('Priority Actions'!$A$5:$A$28)"],
    ["Questions requiring response", "=COUNTA('Backend Questions'!$A$5:$A$18)"],
]
summary["D11"] = "Measure"
summary["E11"] = "Value"
style_header(summary[11][3:5])
for r, (label, formula) in enumerate(facts, start=12):
    summary.cell(r, 4, label)
    summary.cell(r, 5, formula)
    summary.cell(r, 4).border = TABLE_BORDER
    summary.cell(r, 5).border = TABLE_BORDER
    summary.cell(r, 5).alignment = Alignment(horizontal="center")

summary.merge_cells("A17:F17")
summary["A17"] = "Immediate message to the backend team"
summary["A17"].fill = PatternFill("solid", fgColor=BLUE)
summary["A17"].font = Font(name="Aptos", size=11, bold=True, color=WHITE)
messages = [
    "Confirm which backend and database contain the authoritative Nigeria data.",
    "Identify every endpoint that still returns mock, fixed, seed-only or temporary-store data.",
    "Provide KPI-to-source-to-formula lineage for every dashboard figure.",
    "Close provider access to Administration and unrelated institutional records.",
    "Agree an implementation and acceptance plan before further frontend integration patches.",
]
for i, message in enumerate(messages, start=18):
    summary.merge_cells(start_row=i, start_column=1, end_row=i, end_column=6)
    summary.cell(i, 1, f"• {message}")
    summary.cell(i, 1).alignment = Alignment(wrap_text=True, vertical="center")
    summary.cell(i, 1).fill = PatternFill("solid", fgColor=LIGHT_GREY if i % 2 == 0 else WHITE)
    summary.cell(i, 1).border = Border(bottom=THIN)

summary.merge_cells("A25:F25")
summary["A25"] = "Workbook use"
summary["A25"].fill = PatternFill("solid", fgColor=TEAL)
summary["A25"].font = Font(name="Aptos", size=11, bold=True, color=WHITE)
summary.merge_cells("A26:F28")
summary["A26"] = (
    "Use Module Findings as the issue register. The backend team should enter an owner, written response, target date and status. "
    "Use Priority Actions for delivery tracking, Backend Questions for formal clarifications, and Go-Live Criteria for acceptance."
)
summary["A26"].alignment = Alignment(wrap_text=True, vertical="top")
summary["A26"].fill = PatternFill("solid", fgColor=PALE_BLUE)
summary["A26"].border = TABLE_BORDER

chart = BarChart()
chart.type = "bar"
chart.style = 10
chart.title = "Findings by severity"
chart.height = 7
chart.width = 12
chart.add_data(Reference(summary, min_col=2, min_row=11, max_row=14), titles_from_data=True)
chart.set_categories(Reference(summary, min_col=1, min_row=12, max_row=14))
chart.legend = None
chart.x_axis.title = "Number of findings"
summary.add_chart(chart, "H4")
summary.column_dimensions["A"].width = 25
summary.column_dimensions["B"].width = 14
summary.column_dimensions["C"].width = 4
summary.column_dimensions["D"].width = 34
summary.column_dimensions["E"].width = 14
summary.column_dimensions["F"].width = 20
for col in "GHIJ":
    summary.column_dimensions[col].width = 15
summary.freeze_panes = "A4"
set_print(summary, landscape=True)

# Module Findings
title_band(findings, "Module Findings", "Backend team: complete Owner, Backend Response, Target Date and Status. Severity and priority reflect audit impact.", 14)
headers = ["Finding ID", "Module", "Severity", "Assessment", "Expected Behaviour", "Current PHP Behaviour", "Operational Impact", "Required Backend Action", "Priority", "Owner", "Backend Response", "Target Date", "Status", "Evidence / Link"]
findings.append([])
findings.append(headers)
for row in modules:
    findings.append(row + ["", "", "", "Backend Response Required", "See Evidence Sources and full Markdown audit"])
style_header(findings[4])
style_body(findings, 5, 4 + len(modules), 1, len(headers), 72)
widths = [12, 25, 12, 20, 45, 48, 40, 44, 10, 20, 40, 14, 24, 32]
for idx, width in enumerate(widths, 1):
    findings.column_dimensions[findings.cell(4, idx).column_letter].width = width
findings.freeze_panes = "A5"
findings.auto_filter.ref = f"A4:N{4 + len(modules)}"
add_table(findings, "ModuleFindingsTable", f"A4:N{4 + len(modules)}")
add_status_validation(findings, f"M5:M{4 + len(modules)}")
add_priority_validation(findings, f"I5:I{4 + len(modules)}")
findings.conditional_formatting.add(f"C5:C{4 + len(modules)}", FormulaRule(formula=["$C5=\"Critical\""], fill=PatternFill("solid", fgColor=RED_FILL), font=Font(color=RED, bold=True)))
findings.conditional_formatting.add(f"C5:C{4 + len(modules)}", FormulaRule(formula=["$C5=\"High\""], fill=PatternFill("solid", fgColor=ORANGE_FILL), font=Font(color=ORANGE, bold=True)))
findings.conditional_formatting.add(f"C5:C{4 + len(modules)}", FormulaRule(formula=["$C5=\"Medium\""], fill=PatternFill("solid", fgColor=YELLOW_FILL), font=Font(color=YELLOW, bold=True)))
findings.conditional_formatting.add(f"M5:M{4 + len(modules)}", FormulaRule(formula=["$M5=\"Resolved\""], fill=PatternFill("solid", fgColor=GREEN_FILL), font=Font(color=GREEN, bold=True)))
set_print(findings, landscape=True, repeat_rows="1:4")

# Missing Endpoints
title_band(endpoints, "Missing Frontend Endpoints", "Routes currently called or exposed by the frontend but absent from the deployed PHP route table.", 11)
headers = ["ID", "Method", "Route", "Module", "Expected Behaviour", "Current Impact", "Priority", "Owner", "Backend Response", "Target Date", "Status"]
endpoints.append([]); endpoints.append(headers)
for row in missing_endpoints:
    endpoints.append(row + ["", "", "", "Backend Response Required"])
style_header(endpoints[4]); style_body(endpoints, 5, 4 + len(missing_endpoints), 1, len(headers), 62)
for i, w in enumerate([10, 10, 42, 24, 46, 38, 10, 20, 40, 14, 24], 1): endpoints.column_dimensions[endpoints.cell(4, i).column_letter].width = w
endpoints.freeze_panes = "A5"; add_table(endpoints, "MissingEndpointsTable", f"A4:K{4 + len(missing_endpoints)}")
add_status_validation(endpoints, f"K5:K{4 + len(missing_endpoints)}"); add_priority_validation(endpoints, f"G5:G{4 + len(missing_endpoints)}")
set_print(endpoints, True, "1:4")

# Priority Actions
title_band(priority, "Prioritized Backend Action Plan", "P0 = immediate security/truthfulness risk; P1 = data spine; P2 = workflow completion; P3 = acceptance and governance.", 11)
headers = ["Action ID", "Priority", "Action", "Area", "Acceptance Evidence", "Suggested Owner", "Assigned Owner", "Backend Response / Plan", "Target Date", "Status", "% Complete"]
priority.append([]); priority.append(headers)
for row in actions:
    priority.append(row + ["", "", "", "Backend Response Required", 0])
style_header(priority[4]); style_body(priority, 5, 4 + len(actions), 1, len(headers), 58)
for i, w in enumerate([11, 10, 54, 24, 48, 20, 20, 42, 14, 24, 12], 1): priority.column_dimensions[priority.cell(4, i).column_letter].width = w
priority.freeze_panes = "A5"; add_table(priority, "PriorityActionsTable", f"A4:K{4 + len(actions)}")
add_status_validation(priority, f"J5:J{4 + len(actions)}"); add_priority_validation(priority, f"B5:B{4 + len(actions)}")
for cell in priority[f"K5:K{4 + len(actions)}"][0]: pass
for row in range(5, 5 + len(actions)):
    priority.cell(row, 11).number_format = "0%"
priority.conditional_formatting.add(f"K5:K{4 + len(actions)}", FormulaRule(formula=["$K5=1"], fill=PatternFill("solid", fgColor=GREEN_FILL), font=Font(color=GREEN, bold=True)))
set_print(priority, True, "1:4")

# Questions
title_band(questions_ws, "Questions Requiring Written Backend Response", "Responses should include links to code, database mappings, API documentation, tests or environment evidence—not only verbal confirmation.", 8)
headers = ["No.", "Question", "Topic", "Backend Response", "Evidence / Link", "Owner", "Response Due", "Status"]
questions_ws.append([]); questions_ws.append(headers)
for row in questions:
    questions_ws.append(row + ["", "", "", "", "Backend Response Required"])
style_header(questions_ws[4]); style_body(questions_ws, 5, 4 + len(questions), 1, len(headers), 64)
for i, w in enumerate([8, 82, 28, 55, 42, 20, 14, 24], 1): questions_ws.column_dimensions[questions_ws.cell(4, i).column_letter].width = w
questions_ws.freeze_panes = "A5"; add_table(questions_ws, "BackendQuestionsTable", f"A4:H{4 + len(questions)}")
add_status_validation(questions_ws, f"H5:H{4 + len(questions)}"); set_print(questions_ws, True, "1:4")

# Go-live criteria
title_band(criteria_ws, "Go-Live Acceptance Criteria", "A module should not be called live-data-ready until all applicable gates pass without frontend mocks.", 9)
headers = ["Gate ID", "Area", "Required Condition", "Verification Method", "Current Audit Result", "Backend Evidence", "QA Owner", "Retest Date", "Final Status"]
criteria_ws.append([]); criteria_ws.append(headers)
for row in criteria:
    criteria_ws.append(row + ["", "", "", "Open"])
style_header(criteria_ws[4]); style_body(criteria_ws, 5, 4 + len(criteria), 1, len(headers), 64)
for i, w in enumerate([10, 24, 52, 48, 20, 45, 20, 14, 18], 1): criteria_ws.column_dimensions[criteria_ws.cell(4, i).column_letter].width = w
criteria_ws.freeze_panes = "A5"; add_table(criteria_ws, "GoLiveCriteriaTable", f"A4:I{4 + len(criteria)}")
status_gate = DataValidation(type="list", formula1='"Open,Passed,Failed,Not Applicable"', allow_blank=True)
criteria_ws.add_data_validation(status_gate); status_gate.add(f"I5:I{4 + len(criteria)}")
criteria_ws.conditional_formatting.add(f"E5:E{4 + len(criteria)}", FormulaRule(formula=["$E5=\"Failed\""], fill=PatternFill("solid", fgColor=RED_FILL), font=Font(color=RED, bold=True)))
criteria_ws.conditional_formatting.add(f"I5:I{4 + len(criteria)}", FormulaRule(formula=["$I5=\"Passed\""], fill=PatternFill("solid", fgColor=GREEN_FILL), font=Font(color=GREEN, bold=True)))
set_print(criteria_ws, True, "1:4")

# Evidence
title_band(sources_ws, "Evidence Sources", "Project documents and codebases reviewed for this audit. Paths are relative to the frontend workspace unless noted.", 5)
headers = ["Source ID", "Type", "File / Location", "Audit Use", "Backend Comment"]
sources_ws.append([]); sources_ws.append(headers)
for row in sources:
    sources_ws.append(row + [""])
style_header(sources_ws[4]); style_body(sources_ws, 5, 4 + len(sources), 1, len(headers), 54)
for i, w in enumerate([11, 30, 72, 70, 42], 1): sources_ws.column_dimensions[sources_ws.cell(4, i).column_letter].width = w
sources_ws.freeze_panes = "A5"; add_table(sources_ws, "EvidenceSourcesTable", f"A4:E{4 + len(sources)}")
set_print(sources_ws, True, "1:4")

# Common styling and workbook properties
for ws in wb.worksheets:
    ws.sheet_view.zoomScale = 85
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.oddHeader.center.text = "M300 Monitoring Platform — Backend Logic Audit"
    ws.oddFooter.left.text = "Prepared 25 August 2026"
    ws.oddFooter.right.text = "Page &P of &N"

wb.calculation.fullCalcOnLoad = True
wb.calculation.forceFullCalc = True
wb.calculation.calcMode = "auto"
wb.properties.title = "M300 Full Dashboard Backend Logic Audit"
wb.properties.subject = "Backend implementation, logic, security and acceptance findings"
wb.properties.creator = "EMRC dashboard audit"
wb.properties.description = "Full-dashboard audit converted into a backend response and remediation tracker."

OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
wb.save(OUTPUT_PATH)

# Structural verification.
check = load_workbook(OUTPUT_PATH, data_only=False)
expected_sheets = ["Executive Summary", "Module Findings", "Missing Endpoints", "Priority Actions", "Backend Questions", "Go-Live Criteria", "Evidence Sources"]
assert check.sheetnames == expected_sheets
assert check["Module Findings"].max_row == 19
assert len(check["Module Findings"].tables) == 1
assert len(check["Missing Endpoints"].tables) == 1
assert len(check["Priority Actions"].tables) == 1
assert len(check["Backend Questions"].tables) == 1
assert len(check["Go-Live Criteria"].tables) == 1
assert len(check["Evidence Sources"].tables) == 1
assert len(check["Executive Summary"]._charts) == 1
assert check["Executive Summary"]["B12"].value.startswith("=COUNTIF")
for sheet in check.worksheets:
    for row in sheet.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and any(err in cell.value for err in ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A")):
                raise AssertionError(f"Formula error marker in {sheet.title}!{cell.coordinate}: {cell.value}")

print(OUTPUT_PATH)
print(f"sheets={len(check.sheetnames)} findings={len(modules)} actions={len(actions)} questions={len(questions)}")
