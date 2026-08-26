from pathlib import Path
import json
import sys

import openpyxl
from docx import Document
from pptx import Presentation
from pypdf import PdfReader

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
ROOT = Path(__file__).resolve().parents[2]


def clean(value):
    return " ".join(str(value).split()) if value is not None else ""


def emit(label, value, limit=900):
    text = clean(value)
    if text:
        print(f"{label}: {text[:limit]}")


docx_path = next(ROOT.glob("docs/EMRC - Inception Report*.docx"))
print(f"\nDOCX {docx_path.name}")
doc = Document(docx_path)
for index, paragraph in enumerate(doc.paragraphs, 1):
    text = clean(paragraph.text)
    if text and (paragraph.style.name.startswith("Heading") or any(
        key in text.lower() for key in (
            "dashboard", "workflow", "validation", "role", "source", "formula",
            "audit", "evidence", "upload", "report", "pillar", "bottleneck"
        )
    )):
        emit(f"P{index} [{paragraph.style.name}]", text)
for table_index, table in enumerate(doc.tables, 1):
    print(f"TABLE {table_index} rows={len(table.rows)} cols={len(table.columns)}")
    for row_index, row in enumerate(table.rows[:15], 1):
        cells = [clean(cell.text) for cell in row.cells]
        if any(cells):
            emit(f"T{table_index}R{row_index}", json.dumps(cells, ensure_ascii=False), 1300)


for xlsx_path in [
    ROOT / "docs" / "EMRC - M300 KPI Matrix v2-0 2026-06-29.xlsx",
    ROOT / "docs" / "source" / "EMRC - M300  Dashboard Page Specifications v0-3 2026-07-17.xlsx",
]:
    print(f"\nXLSX {xlsx_path.name}")
    workbook = openpyxl.load_workbook(xlsx_path, data_only=False, read_only=True)
    print("SHEETS", workbook.sheetnames)
    for worksheet in workbook.worksheets:
        print(f"SHEET {worksheet.title} rows={worksheet.max_row} cols={worksheet.max_column}")
        shown = 0
        for row in worksheet.iter_rows(values_only=True):
            values = [clean(value) for value in row]
            while values and not values[-1]:
                values.pop()
            if not any(values):
                continue
            emit("ROW", json.dumps(values, ensure_ascii=False), 1600)
            shown += 1
            if shown >= 20:
                break


pptx_path = ROOT / "docs" / "M300 Monitoring Platform Feedback.pptx"
print(f"\nPPTX {pptx_path.name}")
deck = Presentation(pptx_path)
for slide_index, slide in enumerate(deck.slides, 1):
    texts = []
    for shape in slide.shapes:
        if hasattr(shape, "text") and clean(shape.text):
            texts.append(clean(shape.text))
    notes = ""
    try:
        notes = clean(slide.notes_slide.notes_text_frame.text)
    except Exception:
        pass
    emit(f"SLIDE {slide_index}", " | ".join(texts), 2500)
    emit(f"NOTES {slide_index}", notes, 2500)


pdf_path = ROOT / "docs" / "source" / "EMRC - M300 Dashboard Configuration Wireframe v0-1 2026-06-10.pdf"
print(f"\nPDF {pdf_path.name}")
reader = PdfReader(pdf_path)
print("PAGES", len(reader.pages))
for page_index, page in enumerate(reader.pages, 1):
    emit(f"PAGE {page_index}", page.extract_text() or "", 1800)
